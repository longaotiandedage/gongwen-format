#!/usr/bin/env python3
# -*- coding: utf-8 -*-
r"""公文版面校验：量出实测尺寸，与 GB/T 9704—2012 逐项对照。

用法：
    python scripts/verify_layout.py examples/notice.pdf
    python scripts/verify_layout.py examples/notice.docx

PDF 用 PyMuPDF（pip install pymupdf），DOCX 用 python-docx（pip install python-docx）。
只装了其中一个也能用，缺库时给出安装提示后退出。

容差说明：国标为 ±1 mm，但文字包围盒含字体上下伸部，实测值通常比理论值
大半毫米到一毫米，因此判定容差取 ±2 mm。
"""
import os
import sys

MM = 25.4 / 72.0          # pt -> mm
TOL = 2.0                 # 判定容差（mm）

# 国标目标值（mm）
TARGET = {
    "page_w": 210.0,
    "page_h": 297.0,
    "top": 37.0,
    "bottom": 35.0,
    "left": 28.0,
    "right": 26.0,
    "body_w": 156.0,
    "body_h": 225.0,
}


def mm(v):
    return v * MM


def verdict(actual, target, tol=TOL):
    ok = abs(actual - target) <= tol
    return "OK " if ok else "偏差", ok


def report(rows):
    """rows: [(项目, 实测, 国标, 判定)]"""
    w = max(len(r[0]) for r in rows)
    print(f"{'项目':<{w}}  {'实测(mm)':>9}  {'国标(mm)':>9}  判定")
    print("-" * (w + 32))
    all_ok = True
    for name, a, t, mark in rows:
        print(f"{name:<{w}}  {a:>9.1f}  {t:>9.1f}  {mark}")
        if mark != "OK ":
            all_ok = False
    return all_ok


# ---------------------------------------------------------------- PDF
def check_pdf(path):
    try:
        import fitz
    except ImportError:
        print("缺少 PyMuPDF，请先执行： pip install pymupdf", file=sys.stderr)
        return None

    doc = fitz.open(path)
    page = doc[0]
    W, H = page.rect.width, page.rect.height

    # 收集正文行，排除页码（形如 — 1 — 的纯数字行）
    body = []
    for blk in page.get_text("dict")["blocks"]:
        for ln in blk.get("lines", []):
            text = "".join(s["text"] for s in ln["spans"]).strip()
            if not text:
                continue
            stripped = text.strip("—- ").strip()
            if stripped.isdigit() and mm(ln["bbox"][1]) > TARGET["page_h"] - 45:
                continue                      # 页脚页码，不计入正文
            body.append(ln["bbox"])

    if not body:
        print("PDF 里没有提取到文字，可能是扫描件。", file=sys.stderr)
        return None

    left = mm(min(b[0] for b in body))
    right_m = mm(W - max(b[2] for b in body))
    top = mm(min(b[1] for b in body))
    bottom_m = mm(H - max(b[3] for b in body))
    body_w = mm(max(b[2] for b in body) - min(b[0] for b in body))
    body_h = mm(max(b[3] for b in body) - min(b[1] for b in body))

    rows = []
    for name, actual, key in [
        ("页面宽", mm(W), "page_w"),
        ("页面高", mm(H), "page_h"),
        ("上白边（天头）", top, "top"),
        ("下白边", bottom_m, "bottom"),
        ("左白边（订口）", left, "left"),
        ("右白边", right_m, "right"),
        ("版心宽", body_w, "body_w"),
    ]:
        v, _ = verdict(actual, TARGET[key])
        rows.append((name, actual, TARGET[key], v))

    # 版心高只检验不超出，内容不满属正常
    if body_h > TARGET["body_h"] + TOL:
        rows.append(("版心高", body_h, TARGET["body_h"], "超出"))
    else:
        rows.append(("版心高（内容占用）", body_h, TARGET["body_h"], "OK "))

    print(f"文件：{path}    页数：{doc.page_count}")
    return report(rows)


# ---------------------------------------------------------------- DOCX
def check_docx(path):
    try:
        from docx import Document
        from docx.oxml.ns import qn
    except ImportError:
        print("缺少 python-docx，请先执行： pip install python-docx", file=sys.stderr)
        return None

    d = Document(path)
    s = d.sections[0]
    to_mm = 1.0 / 36000.0

    W, H = s.page_width * to_mm, s.page_height * to_mm
    rows = []
    for name, actual, key in [
        ("页面宽", W, "page_w"),
        ("页面高", H, "page_h"),
        ("上白边（天头）", s.top_margin * to_mm, "top"),
        ("下白边", s.bottom_margin * to_mm, "bottom"),
        ("左白边（订口）", s.left_margin * to_mm, "left"),
        ("右白边", s.right_margin * to_mm, "right"),
        ("版心宽", (s.page_width - s.left_margin - s.right_margin) * to_mm, "body_w"),
        ("版心高", (s.page_height - s.top_margin - s.bottom_margin) * to_mm, "body_h"),
    ]:
        v, _ = verdict(actual, TARGET[key])
        rows.append((name, actual, TARGET[key], v))

    # 抽样检查正文字体与行距
    sample = None
    for p in d.paragraphs:
        if p.runs and p.paragraph_format.line_spacing:
            sample = p
            break
    print(f"文件：{path}    段落数：{len(d.paragraphs)}")
    if sample is not None:
        r = sample.runs[0]
        ea = r._element.rPr.rFonts.get(qn("w:eastAsia")) if r._element.rPr is not None else None
        print(f"抽样段落：字体={r.font.name} / eastAsia={ea} / "
              f"字号={r.font.size.pt if r.font.size else '?'}pt / "
              f"行距={sample.paragraph_format.line_spacing.pt:.1f}pt")
    return report(rows)


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    path = sys.argv[1]
    if not os.path.exists(path):
        print(f"找不到文件：{path}", file=sys.stderr)
        return 2

    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        ok = check_pdf(path)
    elif ext == ".docx":
        ok = check_docx(path)
    else:
        print("只支持 .pdf 与 .docx", file=sys.stderr)
        return 2

    if ok is None:
        return 2
    print()
    print("结论：" + ("全部落在国标容差内" if ok else "存在偏差，见上表"))
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
