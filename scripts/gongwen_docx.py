#!/usr/bin/env python3
# -*- coding: utf-8 -*-
r"""公文格式 DOCX 生成器（GB/T 9704—2012）。

可当库用：
    from gongwen_docx import GongwenDoc
    doc = GongwenDoc()
    doc.org("××大学文件")
    doc.docno("教字〔2026〕15 号")
    doc.title("关于××的通知")
    doc.to("各教学单位：")
    doc.body("正文……")
    doc.h1("一、总体要求")
    doc.save("out.docx")

也可直接跑，生成内置样例：
    python scripts/gongwen_docx.py out.docx
"""
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

# 字体：与 gongwen.sty 的回退链保持一致。可按平台改这里。
FANGSONG = "仿宋"
HEITI = "黑体"
KAITI = "楷体"
SONGTI = "宋体"
BIAOSONG = "华文中宋"        # 替代方正小标宋简体

RED = RGBColor(0xD4, 0x00, 0x00)

BODY_PT = 16                 # 三号
LINE_PT = 28                 # 固定行距 28 磅


class GongwenDoc:
    """按国标排版的公文文档。"""

    def __init__(self, typography=None):
        self.doc = Document()
        self.typo = {
            "body": FANGSONG, "h1": HEITI, "h2": KAITI,
            "page": SONGTI, "title": BIAOSONG,
        }
        if typography:
            self.typo.update(typography)
        self._setup_page()
        self._setup_page_numbers()

    # ------------------------------------------------------------ 页面
    def _setup_page(self):
        s = self.doc.sections[0]
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.top_margin, s.bottom_margin = Mm(37), Mm(35)
        s.left_margin, s.right_margin = Mm(28), Mm(26)
        s.header_distance = Mm(15)
        s.footer_distance = Mm(8)

    def _setup_page_numbers(self):
        """奇偶页不同页码：单页居右、双页居左。"""
        settings = self.doc.settings.element
        if settings.find(qn("w:evenAndOddHeaders")) is None:
            settings.append(OxmlElement("w:evenAndOddHeaders"))

        s = self.doc.sections[0]
        self._add_page_field(s.footer.paragraphs[0], WD_ALIGN_PARAGRAPH.RIGHT)
        if hasattr(s, "even_page_footer"):
            s.even_page_footer.is_linked_to_previous = False
            self._add_page_field(s.even_page_footer.paragraphs[0],
                                 WD_ALIGN_PARAGRAPH.LEFT)

    def _add_page_field(self, p, align):
        p.alignment = align
        self._style_run(p.add_run("— "), self.typo["page"], 14)
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), " PAGE ")
        inner = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        rf = OxmlElement("w:rFonts")
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
            rf.set(qn(attr), self.typo["page"])
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "28")            # 半磅，14 磅
        rpr.append(rf)
        rpr.append(sz)
        t = OxmlElement("w:t")
        t.text = "1"
        inner.append(rpr)
        inner.append(t)
        fld.append(inner)
        p._p.append(fld)
        self._style_run(p.add_run(" —"), self.typo["page"], 14)

    # ------------------------------------------------------------ 底层
    @staticmethod
    def _style_run(run, font, size_pt, bold=False, color=None):
        run.font.name = font
        run.font.size = Pt(size_pt)
        run.font.bold = bold
        if color is not None:
            run.font.color.rgb = color
        rpr = run._element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = OxmlElement("w:rFonts")
            rpr.append(rf)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
            rf.set(qn(attr), font)

    def _para(self, text, font=None, size=BODY_PT, bold=False, align=None,
              indent_first=True, line_pt=LINE_PT, color=None,
              before=0, after=0, left_indent_pt=None, right_indent_pt=None):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = Pt(line_pt)
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.first_line_indent = Pt(size * 2) if indent_first else Pt(0)
        if left_indent_pt is not None:
            pf.left_indent = Pt(left_indent_pt)
        if right_indent_pt is not None:
            pf.right_indent = Pt(right_indent_pt)
        if align is not None:
            p.alignment = align
        self._style_run(p.add_run(text), font or self.typo["body"], size,
                        bold, color)
        return p

    # ------------------------------------------------------------ 部件
    def org(self, text):
        """发文机关标志：红色小标宋，居中。"""
        return self._para(text, self.typo["title"], 36,
                          align=WD_ALIGN_PARAGRAPH.CENTER, indent_first=False,
                          line_pt=44, color=RED, after=4)

    def docno(self, text):
        """发文字号，并自动加红色分隔线。"""
        self._para(text, self.typo["page"], BODY_PT,
                   align=WD_ALIGN_PARAGRAPH.CENTER, indent_first=False)
        return self._red_rule()

    def _red_rule(self, before=2, after=16):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = Pt(1)
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.first_line_indent = Pt(0)
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")          # 12 个 1/8 磅 = 1.5 磅
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "D40000")
        pBdr.append(bottom)
        pPr.append(pBdr)
        return p

    def title(self, text):
        """文件标题：2 号小标宋，居中。"""
        return self._para(text, self.typo["title"], 22,
                          align=WD_ALIGN_PARAGRAPH.CENTER, indent_first=False,
                          line_pt=34, after=8)

    def to(self, text):
        """主送机关：顶格。"""
        return self._para(text, indent_first=False)

    def body(self, text):
        """正文段落：首行缩进 2 字。"""
        return self._para(text)

    def h1(self, text):
        """一级标题：3 号黑体。"""
        return self._para(text, self.typo["h1"])

    def h2(self, text):
        """二级标题：3 号楷体。"""
        return self._para(text, self.typo["h2"])

    def h3(self, text):
        """三级标题：3 号仿宋加粗。"""
        return self._para(text, bold=True)

    def attach(self, items):
        """附件说明：首行左空二字，续行左空六字。"""
        if isinstance(items, str):
            items = [items]
        first = items[0]
        self._para("附件：" + first, indent_first=False, before=16,
                   left_indent_pt=2 * BODY_PT)
        for it in items[1:]:
            self._para(it, indent_first=False, left_indent_pt=6 * BODY_PT)

    def sign(self, org_name, date):
        """落款：署名右空二字，日期右空四字。"""
        self._para(org_name, align=WD_ALIGN_PARAGRAPH.RIGHT, indent_first=False,
                   before=24, right_indent_pt=2 * BODY_PT)
        return self._para(date, align=WD_ALIGN_PARAGRAPH.RIGHT,
                          indent_first=False, right_indent_pt=4 * BODY_PT)

    def save(self, path):
        self.doc.save(path)
        return path


# ---------------------------------------------------------------- 样例
DEMO = [
    ("org", "华中科技大学教务处文件"),
    ("docno", "教字〔2026〕15 号"),
    ("title", "关于开展数值方法课程资源建设的通知"),
    ("to", "各教学单位："),
    ("body", "为加强数值计算方法课程建设，提升学生运用计算机求解数学问题的能力，"
             "经研究决定开展课程资源建设工作。现将有关事项通知如下。"),
    ("h1", "一、建设目标"),
    ("h2", "（一）总体要求"),
    ("body", "以培养学生数值计算能力为核心，建设一套内容准确、体例统一、"
             "可直接运行的中文教学资源。"),
    ("h2", "（二）具体目标"),
    ("h3", "1. 讲义建设"),
    ("body", "编写中文讲义一部，正文不少于四十万字，配图不少于一百五十幅。"),
    ("h3", "2. 习题配套"),
    ("body", "每章配习题六至十道，其中二至三道给出提示。"),
    ("h1", "二、工作要求"),
    ("body", "各教学单位要高度重视，统筹安排人员，保证编写时间。"),
]


def build_demo(path):
    d = GongwenDoc()
    for kind, text in DEMO:
        getattr(d, kind)(text)
    d.attach(["1. 课程资源建设任务分工表", "2. 讲义编写体例说明"])
    d.sign("华中科技大学教务处", "二〇二六年十月二十六日")
    return d.save(path)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "gongwen-demo.docx"
    p = build_demo(out)
    print(f"已生成：{p}")
